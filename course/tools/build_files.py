# -*- coding: utf-8 -*-
"""
בונה את קבצי התרגול (course/files) ואת קבצי הפתרון (course/files/solutions)
לקורס "מסלול חשבת". הנתונים בדויים לחלוטין. הרצה: python3 course/tools/build_files.py
"""
import random
import datetime as dt
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1] / "files"
SOL = ROOT / "solutions"
ROOT.mkdir(exist_ok=True)
SOL.mkdir(exist_ok=True)

HEAD_FILL = PatternFill("solid", fgColor="1F4E5A")
HEAD_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")   # צהוב = כאן ממלאים
SOL_FILL = PatternFill("solid", fgColor="E2F0D9")     # ירוק = פתרון
TITLE_FONT = Font(bold=True, size=14, color="1F4E5A")
THIN = Side(style="thin", color="B7C4C8")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = '#,##0;[Red]-#,##0'
NUM2 = '#,##0.00;[Red]-#,##0.00'
PCT = '0.0%'
DATE = 'dd/mm/yyyy'
VAT = 0.18  # שיעור מע"מ בישראל מ-1.1.2025


def new_sheet(wb, title, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.sheet_view.rightToLeft = True
    return ws


def instructions(ws, title, lines):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.column_dimensions["A"].width = 110
    r = 3
    for line in lines:
        ws.cell(row=r, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
        if line.startswith("שלב") or line.startswith("משימה") or line.startswith("מקרא"):
            ws.cell(row=r, column=1).font = Font(bold=True)
        r += 1
    return r


def table(ws, top, headers, rows, widths=None, formats=None, input_cols=(), start_col=1):
    for j, h in enumerate(headers):
        c = ws.cell(row=top, column=start_col + j, value=h)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            c = ws.cell(row=top + i, column=start_col + j, value=v)
            c.border = BOX
            if formats and formats.get(j):
                c.number_format = formats[j]
            if j in input_cols:
                c.fill = INPUT_FILL
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(start_col + j)].width = w
    ws.freeze_panes = ws.cell(row=top + 1, column=start_col)
    return top + len(rows)


def mark(ws, cells, fill):
    for ref in cells:
        ws[ref].fill = fill


LEGEND = "מקרא צבעים: כותרת כהה = נתון קבוע. תא צהוב = כאן את ממלאת. בקובץ הפתרון תאים ירוקים = התשובה."

# ---------------------------------------------------------------------------
# שלב 2: פקודות יומן
# ---------------------------------------------------------------------------
TX = [
    (1, "02/01/2026", "בעלת המניות הפקידה 100,000 ₪ בחשבון הבנק של החברה כהון ראשוני",
     [("בנק", 100000, 0), ("הון מניות", 0, 100000)]),
    (2, "03/01/2026", "קנינו מכונת אספרסו ב-30,000 ₪ + מע\"מ, שולם בהעברה בנקאית",
     [("ציוד (רכוש קבוע)", 30000, 0), ("מע\"מ תשומות", 5400, 0), ("בנק", 0, 35400)]),
    (3, "05/01/2026", "קנינו פולי קפה מספק ב-5,000 ₪ + מע\"מ, נשלם בעוד 30 יום (באשראי)",
     [("קניות מלאי", 5000, 0), ("מע\"מ תשומות", 900, 0), ("ספקים", 0, 5900)]),
    (4, "31/01/2026", "מכירות בקופה בחודש ינואר: 23,600 ₪ כולל מע\"מ, נכנסו לבנק",
     [("בנק", 23600, 0), ("הכנסות ממכירות", 0, 20000), ("מע\"מ עסקאות", 0, 3600)]),
    (5, "01/01/2026", "שילמנו שכירות לחודש ינואר: 8,000 ₪ + מע\"מ",
     [("הוצאות שכירות", 8000, 0), ("מע\"מ תשומות", 1440, 0), ("בנק", 0, 9440)]),
    (6, "04/02/2026", "שילמנו לספק הפולים את החוב מעסקה 3",
     [("ספקים", 5900, 0), ("בנק", 0, 5900)]),
    (7, "20/01/2026", "הוצאנו חשבונית ללקוח עסקי על אירוע: 10,000 ₪ + מע\"מ, ישלם בעוד 30 יום",
     [("לקוחות", 11800, 0), ("הכנסות ממכירות", 0, 10000), ("מע\"מ עסקאות", 0, 1800)]),
    (8, "10/01/2026", "קיבלנו הלוואה מהבנק: 50,000 ₪",
     [("בנק", 50000, 0), ("הלוואה מבנק", 0, 50000)]),
    (9, "31/01/2026", "שילמנו משכורות לעובדים: 12,000 ₪ (בפשטות, בלי ניכויים)",
     [("הוצאות שכר", 12000, 0), ("בנק", 0, 12000)]),
    (10, "25/01/2026", "חשבון חשמל: 1,180 ₪ כולל מע\"מ, שולם מהבנק",
     [("הוצאות חשמל", 1000, 0), ("מע\"מ תשומות", 180, 0), ("בנק", 0, 1180)]),
    (11, "19/02/2026", "הלקוח העסקי מעסקה 7 שילם את כל החוב",
     [("בנק", 11800, 0), ("לקוחות", 0, 11800)]),
    (12, "28/01/2026", "קיבלנו חשבונית על פרסום באינסטגרם: 2,000 ₪ + מע\"מ, עוד לא שילמנו",
     [("הוצאות פרסום", 2000, 0), ("מע\"מ תשומות", 360, 0), ("ספקים", 0, 2360)]),
]
ACCOUNTS_01 = ["בנק", "לקוחות", "קניות מלאי", "ציוד (רכוש קבוע)", "מע\"מ תשומות", "ספקים",
               "מע\"מ עסקאות", "הלוואה מבנק", "הון מניות", "הכנסות ממכירות", "הוצאות שכירות",
               "הוצאות שכר", "הוצאות חשמל", "הוצאות פרסום"]
ACC_TYPE_01 = {"בנק": "נכס", "לקוחות": "נכס", "קניות מלאי": "הוצאה", "ציוד (רכוש קבוע)": "נכס",
               "מע\"מ תשומות": "נכס", "ספקים": "התחייבות", "מע\"מ עסקאות": "התחייבות",
               "הלוואה מבנק": "התחייבות", "הון מניות": "הון", "הכנסות ממכירות": "הכנסה",
               "הוצאות שכירות": "הוצאה", "הוצאות שכר": "הוצאה", "הוצאות חשמל": "הוצאה",
               "הוצאות פרסום": "הוצאה"}


def build_01():
    for solution in (False, True):
        wb = Workbook()
        ws = new_sheet(wb, "הוראות", first=True)
        instructions(ws, "שלב 2 | פקודות יומן | קפה נחת בע\"מ", [
            "הסיפור: פתחת בית קפה קטן בשם \"קפה נחת בע\"מ\". את החשבת. כל כסף שזז צריך להירשם.",
            LEGEND,
            "",
            "שלב 1: פתחי את הגיליון \"עסקאות\" וקראי כל עסקה. שאלי את עצמך: מה נכנס? מה יצא?",
            "שלב 2: פתחי את הגיליון \"רשימת חשבונות\". מכאן בוחרים שמות חשבונות. לא ממציאים שמות חדשים.",
            "שלב 3: בגיליון \"פקודות יומן\" רשמי לכל עסקה שורה אחת או יותר: חשבון, חובה, זכות.",
            "שלב 4: בדיקה. בכל פקודה סכום החובה חייב להיות שווה לסכום הזכות. עמודת הבדיקה צריכה להראות 0.",
            "שלב 5: בגיליון \"מאזן בוחן\" סכמי לכל חשבון את כל החובה ואת כל הזכות (רמז: SUMIFS).",
            "שלב 6: בדיקה אחרונה. סך כל החובה במאזן הבוחן = סך כל הזכות. אם לא, יש טעות. חפשי אותה.",
            "",
            "טיפ ADHD: עשי 3 עסקאות, קומי לשתות מים, עוד 3 עסקאות. 12 עסקאות = 4 סבבים קצרים.",
            "טיפ: מע\"מ בישראל הוא 18%. מחיר \"כולל מע\"מ\" מחלקים ב-1.18 כדי לקבל את המחיר בלי מע\"מ.",
        ])
        ws2 = new_sheet(wb, "עסקאות")
        table(ws2, 1, ["מס' עסקה", "תאריך", "מה קרה"], [(t[0], t[1], t[2]) for t in TX], [10, 13, 90])

        ws3 = new_sheet(wb, "רשימת חשבונות")
        table(ws3, 1, ["שם חשבון", "סוג"], [(a, ACC_TYPE_01[a]) for a in ACCOUNTS_01], [24, 14])

        ws4 = new_sheet(wb, "פקודות יומן")
        rows = []
        if solution:
            for t in TX:
                for acc, d, c in t[3]:
                    rows.append([t[0], acc, d or None, c or None])
        else:
            for t in TX:
                for _ in range(max(2, len(t[3]))):
                    rows.append([t[0], None, None, None])
        last = table(ws4, 1, ["מס' עסקה", "חשבון", "חובה", "זכות"], rows, [10, 24, 14, 14],
                     {2: NUM, 3: NUM}, input_cols=() if solution else (1, 2, 3))
        ws4.cell(row=1, column=6, value="מס' עסקה").font = Font(bold=True)
        ws4.cell(row=1, column=7, value="בדיקה: חובה פחות זכות (חייב 0)").font = Font(bold=True)
        ws4.column_dimensions["G"].width = 30
        for i, t in enumerate(TX, start=2):
            ws4.cell(row=i, column=6, value=t[0])
            c = ws4.cell(row=i, column=7,
                         value=f"=SUMIFS($C$2:$C${last},$A$2:$A${last},F{i})-SUMIFS($D$2:$D${last},$A$2:$A${last},F{i})")
            c.number_format = NUM

        ws5 = new_sheet(wb, "מאזן בוחן")
        tb_rows = []
        for i, a in enumerate(ACCOUNTS_01, start=2):
            if solution:
                d = f"=SUMIFS('פקודות יומן'!$C:$C,'פקודות יומן'!$B:$B,A{i})"
                c = f"=SUMIFS('פקודות יומן'!$D:$D,'פקודות יומן'!$B:$B,A{i})"
            else:
                d = c = None
            tb_rows.append([a, ACC_TYPE_01[a], d, c, f"=C{i}-D{i}"])
        n = len(ACCOUNTS_01) + 1
        table(ws5, 1, ["חשבון", "סוג", "סה\"כ חובה", "סה\"כ זכות", "יתרה (חובה פלוס / זכות מינוס)"],
              tb_rows, [24, 12, 14, 14, 30], {2: NUM, 3: NUM, 4: NUM},
              input_cols=() if solution else (2, 3))
        ws5.cell(row=n + 1, column=1, value="סה\"כ").font = Font(bold=True)
        for col in "CDE":
            cc = ws5[f"{col}{n + 1}"]
            cc.value = f"=SUM({col}2:{col}{n})"
            cc.number_format = NUM
            cc.font = Font(bold=True)
        ws5.cell(row=n + 3, column=1, value="אם תא E האחרון מראה 0, המאזן מאוזן. כל הכבוד!")
        if solution:
            for r in range(2, n + 1):
                ws5[f"C{r}"].fill = SOL_FILL
                ws5[f"D{r}"].fill = SOL_FILL
        name = "L02_journal_entries_solution.xlsx" if solution else "L02_journal_entries.xlsx"
        wb.save((SOL if solution else ROOT) / name)


# ---------------------------------------------------------------------------
# שלב 3: מאזן בוחן לדוחות כספיים
# ---------------------------------------------------------------------------
TB02 = [
    # קוד, שם, מחלקה בדוח, יתרה (חובה חיובי / זכות שלילי)
    ("1010", "מזומנים ושווי מזומנים", "מאזן: נכסים שוטפים", 184500),
    ("1100", "לקוחות", "מאזן: נכסים שוטפים", 612000),
    ("1200", "מלאי", "מאזן: נכסים שוטפים", 455000),
    ("1300", "הוצאות מראש", "מאזן: נכסים שוטפים", 36000),
    ("1500", "רכוש קבוע - עלות", "מאזן: נכסים לא שוטפים", 1250000),
    ("1510", "רכוש קבוע - פחת נצבר", "מאזן: נכסים לא שוטפים", -410000),
    ("2000", "ספקים ונותני שירותים", "מאזן: התחייבויות שוטפות", -388000),
    ("2100", "מוסדות (מע\"מ, ניכויים, ביטוח לאומי)", "מאזן: התחייבויות שוטפות", -97500),
    ("2200", "הוצאות לשלם", "מאזן: התחייבויות שוטפות", -64000),
    ("2300", "חלויות שוטפות של הלוואה", "מאזן: התחייבויות שוטפות", -120000),
    ("2500", "הלוואה מבנק לזמן ארוך", "מאזן: התחייבויות לא שוטפות", -480000),
    ("3000", "הון מניות", "מאזן: הון עצמי", -100000),
    ("3100", "עודפים (רווחים מצטברים) לתחילת שנה", "מאזן: הון עצמי", None),  # מחושב
    ("4000", "הכנסות ממכירות", "רווח והפסד: הכנסות", -4850000),
    ("5000", "עלות המכר", "רווח והפסד: עלות המכר", 2910000),
    ("6100", "שכר עבודה ונלוות", "רווח והפסד: הוצאות תפעול", 780000),
    ("6200", "שכירות ואחזקה", "רווח והפסד: הוצאות תפעול", 216000),
    ("6300", "פחת", "רווח והפסד: הוצאות תפעול", 125000),
    ("6400", "שיווק ופרסום", "רווח והפסד: הוצאות תפעול", 184000),
    ("6500", "הוצאות משרד", "רווח והפסד: הוצאות תפעול", 62000),
    ("6600", "שירותים מקצועיים (רו\"ח, עו\"ד)", "רווח והפסד: הוצאות תפעול", 48000),
    ("7000", "הוצאות מימון", "רווח והפסד: מימון", 41000),
    ("7100", "הכנסות מימון", "רווח והפסד: מימון", -6500),
]
_plug = -sum(x[3] for x in TB02 if x[3] is not None)
TB02 = [(a, b, c, _plug if d is None else d) for a, b, c, d in TB02]


def build_02():
    rev = 4850000
    cogs = 2910000
    gross = rev - cogs
    opex = 780000 + 216000 + 125000 + 184000 + 62000 + 48000
    op = gross - opex
    fin = 41000 - 6500
    pbt = op - fin
    tax = round(pbt * 0.23)
    net = pbt - tax
    for solution in (False, True):
        wb = Workbook()
        ws = new_sheet(wb, "הוראות", first=True)
        instructions(ws, "שלב 3 | ממאזן בוחן לדוחות כספיים | גלבוע אספקה בע\"מ | שנת 2025", [
            "הסיפור: רואה החשבון המבקר מבקש ממך טיוטת דוחות כספיים לשנת 2025. יש לך רק מאזן בוחן.",
            LEGEND,
            "",
            "שלב 1: בגיליון \"מאזן בוחן\" בדקי שסך העמודה \"יתרה\" שווה 0. זה אומר שהמאזן מאוזן.",
            "שלב 2: שימי לב לעמודה \"מקום בדוח\". כל חשבון הולך לדוח רווח והפסד או למאזן.",
            "שלב 3: בגיליון \"רווח והפסד\" מלאי את התאים הצהובים. זכרי: הכנסות רשומות במינוס (זכות). בדוח מציגים אותן בפלוס.",
            "שלב 4: חשבי מס חברות 23% על הרווח לפני מס. (בפשטות, בלי התאמות למס.)",
            "שלב 5: בגיליון \"מאזן\" מלאי נכסים, התחייבויות והון. את המס שחישבת הוסיפי כהתחייבות \"מס לשלם\".",
            "שלב 6: עודפים לסוף שנה = עודפים לתחילת שנה + רווח נקי של השנה.",
            "שלב 7: בדיקת הזהב: סך נכסים = סך התחייבויות + הון. אם לא, חזרי לשלב 3.",
            "",
            "שאלת חשיבה (תשובה בקובץ הפתרון): מה שיעור הרווח הגולמי? מה היחס השוטף? האם החברה נראית בריאה?",
        ])
        ws2 = new_sheet(wb, "מאזן בוחן")
        rows = [list(x) for x in TB02]
        last = table(ws2, 1, ["קוד חשבון", "שם חשבון", "מקום בדוח", "יתרה (חובה + / זכות -)"], rows,
                     [11, 38, 30, 22], {3: NUM})
        ws2.cell(row=last + 1, column=2, value="בדיקה: סך הכל (חייב להיות 0)").font = Font(bold=True)
        c = ws2.cell(row=last + 1, column=4, value=f"=SUM(D2:D{last})")
        c.number_format, c.font = NUM, Font(bold=True)

        def tb(code):
            for i, x in enumerate(TB02, start=2):
                if x[0] == code:
                    return f"'מאזן בוחן'!D{i}"
            raise KeyError(code)

        ws3 = new_sheet(wb, "רווח והפסד")
        ws3["A1"], ws3["A1"].font = "דוח רווח והפסד לשנה שהסתיימה ב-31.12.2025 (ש\"ח)", TITLE_FONT
        pl = [
            ("הכנסות ממכירות", f"=-{tb('4000')}"),
            ("עלות המכר", f"={tb('5000')}"),
            ("רווח גולמי", "=B3-B4"),
            ("הוצאות תפעול (סכום 6100 עד 6600)",
             "=" + "+".join(tb(c) for c in ["6100", "6200", "6300", "6400", "6500", "6600"])),
            ("רווח תפעולי", "=B5-B6"),
            ("הוצאות מימון, נטו", f"={tb('7000')}+{tb('7100')}"),
            ("רווח לפני מס", "=B7-B8"),
            ("מס חברות 23%", "=ROUND(B9*0.23,0)"),
            ("רווח נקי", "=B9-B10"),
            ("", None),
            ("שיעור רווח גולמי", "=B5/B3"),
            ("שיעור רווח תפעולי", "=B7/B3"),
            ("שיעור רווח נקי", "=B11/B3"),
        ]
        for i, (lbl, f) in enumerate(pl, start=3):
            ws3.cell(row=i, column=1, value=lbl)
            cell = ws3.cell(row=i, column=2, value=f if solution else None)
            cell.number_format = PCT if i >= 13 else NUM
            cell.fill = SOL_FILL if solution else INPUT_FILL
            if lbl in ("רווח גולמי", "רווח תפעולי", "רווח לפני מס", "רווח נקי"):
                ws3.cell(row=i, column=1).font = Font(bold=True)
        ws3.cell(row=12, column=2).fill = PatternFill()
        ws3.column_dimensions["A"].width = 40
        ws3.column_dimensions["B"].width = 18

        ws4 = new_sheet(wb, "מאזן")
        ws4["A1"], ws4["A1"].font = "מאזן ליום 31.12.2025 (ש\"ח)", TITLE_FONT
        bs = [
            ("נכסים שוטפים", None, True),
            ("מזומנים ושווי מזומנים", f"={tb('1010')}", False),
            ("לקוחות", f"={tb('1100')}", False),
            ("מלאי", f"={tb('1200')}", False),
            ("הוצאות מראש", f"={tb('1300')}", False),
            ("סה\"כ נכסים שוטפים", "=SUM(B4:B7)", True),
            ("נכסים לא שוטפים", None, True),
            ("רכוש קבוע, נטו (עלות פחות פחת נצבר)", f"={tb('1500')}+{tb('1510')}", False),
            ("סה\"כ נכסים", "=B8+B10", True),
            ("", None, False),
            ("התחייבויות שוטפות", None, True),
            ("ספקים ונותני שירותים", f"=-{tb('2000')}", False),
            ("מוסדות", f"=-{tb('2100')}", False),
            ("הוצאות לשלם", f"=-{tb('2200')}", False),
            ("חלויות שוטפות של הלוואה", f"=-{tb('2300')}", False),
            ("מס חברות לשלם", "='רווח והפסד'!B10", False),
            ("סה\"כ התחייבויות שוטפות", "=SUM(B14:B18)", True),
            ("התחייבויות לא שוטפות", None, True),
            ("הלוואה מבנק לזמן ארוך", f"=-{tb('2500')}", False),
            ("הון עצמי", None, True),
            ("הון מניות", f"=-{tb('3000')}", False),
            ("עודפים לסוף שנה", f"=-{tb('3100')}+'רווח והפסד'!B11", False),
            ("סה\"כ הון עצמי", "=B23+B24", True),
            ("סה\"כ התחייבויות והון", "=B19+B21+B25", True),
            ("", None, False),
            ("בדיקה: נכסים פחות (התחייבויות + הון). חייב 0", "=B11-B26", True),
            ("יחס שוטף (נכסים שוטפים / התחייבויות שוטפות)", "=B8/B19", True),
        ]
        for i, (lbl, f, bold) in enumerate(bs, start=3):
            ws4.cell(row=i, column=1, value=lbl).font = Font(bold=bold)
            if f is not None:
                cell = ws4.cell(row=i, column=2, value=f if solution else None)
                cell.number_format = '0.00' if "יחס" in lbl else NUM
                cell.fill = SOL_FILL if solution else INPUT_FILL
        ws4.column_dimensions["A"].width = 48
        ws4.column_dimensions["B"].width = 18
        if solution:
            ws5 = new_sheet(wb, "תשובות לשאלת החשיבה")
            instructions(ws5, "תשובות", [
                f"רווח גולמי: {gross:,.0f} ₪, שיעור רווח גולמי {gross / rev:.1%}.",
                f"רווח תפעולי: {op:,.0f} ₪. רווח לפני מס: {pbt:,.0f} ₪. מס: {tax:,.0f} ₪. רווח נקי: {net:,.0f} ₪.",
                "יחס שוטף מעל 1 אומר שיש לחברה יותר נכסים שוטפים מחובות קצרים. זה סימן טוב לנזילות.",
                "שימי לב: לקוחות (612 אלף) גבוה מאוד ביחס למזומן (184.5 אלף). כדאי לבדוק גיול לקוחות (שלב 5).",
                "הערה: בדוחות אמיתיים לפי IFRS יש עוד סעיפים (מסים נדחים, דוח תזרים, ביאורים). כאן זו טיוטה לתרגול.",
            ])
        wb.save((SOL if solution else ROOT) / ("L03_trial_balance_to_fs" + ("_solution" if solution else "") + ".xlsx"))


# ---------------------------------------------------------------------------
# שלב 4: אקסל לחשבת
# ---------------------------------------------------------------------------
def build_03():
    rnd = random.Random(42)
    customers = [
        ("C100", "רשת מזון השרון", "נתניה", 30), ("C101", "סופר גליל", "כרמיאל", 60),
        ("C102", "מרכולית הדרום", "באר שבע", 45), ("C103", "קפה פינתי", "תל אביב", 30),
        ("C104", "מלון הים", "אילת", 60), ("C105", "קייטרינג זהב", "חיפה", 30),
        ("C106", "בית ספר אורט", "ירושלים", 90), ("C107", "מסעדת זית", "ראשון לציון", 30),
        ("C108", "מיני מרקט הכפר", "עפולה", 45), ("C109", "בית אבות נווה", "רחובות", 60),
    ]
    products = [
        ("P01", "פולי קפה 1 ק\"ג", "קפה", 48, 85), ("P02", "קפסולות 100 יח'", "קפה", 60, 110),
        ("P03", "חלב שיבולת 1 ל'", "משקאות", 7, 13), ("P04", "סירופ וניל", "משקאות", 18, 32),
        ("P05", "כוסות נייר 1000", "חד פעמי", 90, 150), ("P06", "מכסים 1000", "חד פעמי", 40, 70),
        ("P07", "עוגיות חמאה ארגז", "מאפים", 55, 95), ("P08", "קרואסון קפוא 50", "מאפים", 120, 199),
    ]
    rows = []
    inv = 50001
    start = dt.date(2026, 1, 1)
    for _ in range(400):
        d = start + dt.timedelta(days=rnd.randint(0, 180))
        cust = rnd.choice(customers)[0]
        prod = rnd.choice(products)
        qty = rnd.randint(1, 40)
        price = prod[4]
        if rnd.random() < 0.08:
            price = round(price * 0.9)  # הנחה
        rows.append([d, inv, cust, prod[0], qty, price])
        inv += 1
    rows.sort(key=lambda r: r[0])
    for i, r in enumerate(rows):
        r[1] = 50001 + i
    # שתלנו: 3 כפילויות ו-2 קודי לקוח שלא קיימים
    dups = [rows[37][:], rows[188][:], rows[301][:]]
    rows.insert(38, dups[0])
    rows.insert(190, dups[1])
    rows.insert(304, dups[2])
    rows[120][2] = "C111"
    rows[250][2] = "C112"
    cmap = {c[0]: c for c in customers}
    pmap = {p[0]: p for p in products}

    for solution in (False, True):
        wb = Workbook()
        ws = new_sheet(wb, "הוראות", first=True)
        instructions(ws, "שלב 4 | אקסל לחשבת | חברת \"אספקת בוקר בע\"מ\" | מכירות ינואר עד יוני 2026", [
            "הסיפור: המנכ\"ל שואל שאלות על המכירות. יש לך קובץ גולמי מהמערכת (ייצוא). התשובות צריכות להיות מדויקות.",
            LEGEND,
            "",
            "משימה 1 (XLOOKUP): בגיליון \"מכירות\" מלאי בעמודה G את שם הלקוח מגיליון \"לקוחות\".",
            "   נוסחה לדוגמה: =XLOOKUP(C2,'לקוחות'!A:A,'לקוחות'!B:B,\"לא נמצא\")",
            "משימה 2: בעמודה H חשבי סכום לפני מע\"מ = כמות × מחיר. בעמודה I: חודש =MONTH(A2).",
            "משימה 3 (עלות ורווח): בעמודה J הביאי עלות ליחידה מגיליון \"מוצרים\" וחשבי בעמודה K רווח גולמי = (מחיר - עלות) × כמות.",
            "משימה 4 (SUMIFS): בגיליון \"שאלות המנכ\"ל\" ענו על כל השאלות בתאים הצהובים.",
            "משימה 5 (טבלת ציר / Pivot): הוסיפי > טבלת ציר. שורות = שם לקוח. עמודות = חודש. ערכים = סכום.",
            "משימה 6 (בקרה): מצאי חשבוניות כפולות. רמז: =COUNTIF(B:B,B2)>1 ואז סינון.",
            "משימה 7 (בקרה): מצאי שורות עם קוד לקוח שלא קיים ברשימת הלקוחות. (משימה 1 כבר תראה \"לא נמצא\".)",
            "",
            "למה זה חשוב? בראיון לחשבת כמעט תמיד יש מבחן אקסל. 80% מהמבחנים בודקים בדיוק את הכלים האלה.",
            "אין לך XLOOKUP (אקסל ישן)? השתמשי ב-INDEX ו-MATCH או ב-VLOOKUP.",
        ])
        ws2 = new_sheet(wb, "מכירות")
        headers = ["תאריך", "מס' חשבונית", "קוד לקוח", "קוד מוצר", "כמות", "מחיר ליחידה",
                   "שם לקוח", "סכום לפני מע\"מ", "חודש", "עלות ליחידה", "רווח גולמי", "כפולה?"]
        data = []
        for i, r in enumerate(rows, start=2):
            extra = [None] * 6
            if solution:
                extra = [f"=IFERROR(INDEX('לקוחות'!$B:$B,MATCH(C{i},'לקוחות'!$A:$A,0)),\"לא נמצא\")",
                         f"=E{i}*F{i}", f"=MONTH(A{i})",
                         f"=INDEX('מוצרים'!$D:$D,MATCH(D{i},'מוצרים'!$A:$A,0))",
                         f"=(F{i}-J{i})*E{i}", f"=IF(COUNTIF($B:$B,B{i})>1,\"כפולה\",\"\")"]
            data.append(r + extra)
        table(ws2, 1, headers, data, [12, 12, 10, 10, 8, 12, 20, 16, 8, 12, 14, 10],
              {0: DATE, 5: NUM, 7: NUM, 9: NUM, 10: NUM},
              input_cols=() if solution else (6, 7, 8, 9, 10, 11))
        ws2.auto_filter.ref = f"A1:L{len(rows) + 1}"
        ws3 = new_sheet(wb, "לקוחות")
        table(ws3, 1, ["קוד לקוח", "שם לקוח", "עיר", "תנאי תשלום (ימים)"], customers, [10, 22, 14, 18])
        ws4 = new_sheet(wb, "מוצרים")
        table(ws4, 1, ["קוד מוצר", "שם מוצר", "קטגוריה", "עלות ליחידה", "מחיר מחירון"], products,
              [10, 22, 12, 12, 12], {3: NUM, 4: NUM})

        # תשובות מחושבות בפייתון
        clean = []
        seen = set()
        for r in rows:
            key = tuple(r)
            clean.append(r)
        total = sum(r[4] * r[5] for r in rows)
        total_clean = total - sum(d[4] * d[5] for d in dups)
        jun = sum(r[4] * r[5] for r in rows if r[0].month == 6)
        by_c = {}
        for r in rows:
            by_c[r[2]] = by_c.get(r[2], 0) + r[4] * r[5]
        top_c = max((k for k in by_c if k in cmap), key=lambda k: by_c[k])
        coffee = sum(r[4] * r[5] for r in rows if pmap[r[3]][2] == "קפה")
        gp = sum((r[5] - pmap[r[3]][3]) * r[4] for r in rows)
        q1 = sum(r[4] * r[5] for r in rows if r[0].month <= 3)
        q2 = sum(r[4] * r[5] for r in rows if 4 <= r[0].month <= 6)

        ws5 = new_sheet(wb, "שאלות המנכ\"ל")
        qs = [
            ("מה סך המכירות (לפני מע\"מ) בכל הקובץ, כולל הכפולות?", "=SUM('מכירות'!H:H)", total),
            ("מה סך המכירות בחודש יוני?", "=SUMIFS('מכירות'!H:H,'מכירות'!I:I,6)", jun),
            ("איזה לקוח (קוד) קנה הכי הרבה? (רמז: טבלת ציר)", None, top_c),
            ("כמה מכרנו בקטגוריית \"קפה\"? (רמז: הוסיפי עמודת קטגוריה)", None, coffee),
            ("מה הרווח הגולמי הכולל?", "=SUM('מכירות'!K:K)", gp),
            ("כמה חשבוניות כפולות מצאת?", "=COUNTIF('מכירות'!L:L,\"כפולה\")/2", 3),
            ("מה סך המכירות אחרי מחיקת הכפולות?", None, total_clean),
            ("בכמה אחוזים השתנו המכירות מרבעון 1 לרבעון 2?", None, (q2 - q1) / q1),
            ("אילו קודי לקוח לא קיימים ברשימה?", None, "C111, C112"),
        ]
        for i, (q, f, ans) in enumerate(qs, start=2):
            ws5.cell(row=i, column=1, value=q).alignment = Alignment(wrap_text=True)
            c = ws5.cell(row=i, column=2, value=(ans if solution else None))
            c.fill = SOL_FILL if solution else INPUT_FILL
            c.number_format = PCT if isinstance(ans, float) and abs(ans) < 5 else NUM
            if solution and f:
                ws5.cell(row=i, column=3, value="נוסחה אפשרית: " + f)
        ws5["A1"], ws5["B1"] = "שאלה", "תשובה"
        for c in ("A1", "B1", "C1"):
            ws5[c].fill, ws5[c].font = HEAD_FILL, HEAD_FONT
        if solution:
            ws5["C1"] = "איך"
        ws5.column_dimensions["A"].width = 60
        ws5.column_dimensions["B"].width = 18
        ws5.column_dimensions["C"].width = 60
        wb.save((SOL if solution else ROOT) / ("L04_excel_for_controllers" + ("_solution" if solution else "") + ".xlsx"))


# ---------------------------------------------------------------------------
# שלב 5: לקוחות, ספקים, מע"מ
# ---------------------------------------------------------------------------
def build_04():
    rnd = random.Random(7)
    cutoff = dt.date(2026, 6, 30)
    names = ["רשת מזון השרון", "סופר גליל", "מרכולית הדרום", "קפה פינתי", "מלון הים", "קייטרינג זהב",
             "בית ספר אורט", "מסעדת זית", "מיני מרקט הכפר", "בית אבות נווה", "חברת אירועים שמש", "מתנ\"ס גבעה"]
    inv_rows = []
    n = 7001
    for _ in range(45):
        cust = rnd.choice(names)
        d = cutoff - dt.timedelta(days=rnd.choice([5, 12, 20, 28, 35, 44, 50, 63, 71, 85, 95, 110, 130, 160, 200]))
        terms = rnd.choice([30, 45, 60])
        due = d + dt.timedelta(days=terms)
        amt = round(rnd.randint(1500, 48000), -1)
        inv_rows.append([n, cust, d, terms, due, amt])
        n += 1
    inv_rows.sort(key=lambda r: r[2])
    for i, r in enumerate(inv_rows):
        r[0] = 7001 + i
    rates = [(0, "לא בפיגור", 0.0), (1, "1-30", 0.01), (31, "31-60", 0.05), (61, "61-90", 0.2), (91, "מעל 90", 0.5)]

    def bucket(days):
        if days <= 0:
            return "לא בפיגור", 0.0
        for lo, nm, rt in reversed(rates):
            if days >= lo:
                return nm, rt

    sales = [
        ("מכירת סחורה", 142000, "עסקאות"), ("מכירת סחורה", 98500, "עסקאות"),
        ("שירות אירוע", 36000, "עסקאות"), ("יצוא ללקוח בחו\"ל (מע\"מ אפס)", 54000, "יצוא"),
        ("מכירת סחורה", 77250, "עסקאות"),
    ]
    purchases = [
        ("קניית מלאי", 118000, "כן"), ("שכירות משרד", 22000, "כן"), ("ייעוץ משפטי", 8500, "כן"),
        ("כיבוד וארוחות לעובדים", 3200, "לא (כיבוד אינו מוכר לקיזוז)"),
        ("מחשבים וציוד", 16400, "כן (ציוד)"), ("דלק לרכב פרטי של מנהל", 2400, "לא (רכב פרטי)"),
        ("פרסום בגוגל", 9800, "כן"),
    ]
    for solution in (False, True):
        wb = Workbook()
        ws = new_sheet(wb, "הוראות", first=True)
        instructions(ws, "שלב 5 | לקוחות, גיול חובות ומע\"מ | אספקת בוקר בע\"מ | 30.6.2026", [
            "הסיפור: סוף יוני. הבנק שואל כמה כסף הלקוחות חייבים ומתי ייכנס. בנוסף צריך להכין את דוח המע\"מ של יוני.",
            LEGEND,
            "",
            "חלק א: גיול לקוחות (Aging)",
            "שלב 1: בגיליון \"חשבוניות פתוחות\" חשבי בעמודה G ימי פיגור = תאריך הדוח (30/06/2026) פחות תאריך פירעון. אם שלילי, אין פיגור.",
            "שלב 2: בעמודה H שייכי כל חשבונית לדלי: לא בפיגור / 1-30 / 31-60 / 61-90 / מעל 90. (רמז: IFS או VLOOKUP עם התאמה משוערת)",
            "שלב 3: בעמודה I חשבי הפרשה לחובות מסופקים לפי טבלת השיעורים בגיליון \"מדיניות הפרשה\".",
            "שלב 4: בגיליון \"סיכום גיול\" סכמי לפי דלי. מי 3 הלקוחות עם הכי הרבה חוב מעל 90 יום? למי תתקשרי קודם?",
            "",
            "חלק ב: דוח מע\"מ חודשי",
            "שלב 5: בגיליון \"מע\"מ יוני\" חשבי מע\"מ עסקאות 18% על המכירות. יצוא: מע\"מ 0%.",
            "שלב 6: חשבי מע\"מ תשומות על הקניות. רק מה שמותר לקיזוז נכנס לחישוב.",
            "שלב 7: מע\"מ לתשלום = מע\"מ עסקאות פחות מע\"מ תשומות מותר. אם יוצא שלילי, זה החזר.",
            "",
            "הערה: הכללים כאן פשוטים לצורכי לימוד. בפועל יש כללים מפורטים (למשל רכב, כיבוד, חשבוניות עם מספר הקצאה). תמיד בודקים מול יועץ המס.",
        ])
        ws2 = new_sheet(wb, "חשבוניות פתוחות")
        rows = []
        for i, r in enumerate(inv_rows, start=2):
            if solution:
                ext = [f"=MAX(0,DATE(2026,6,30)-E{i})",
                       f"=IF(G{i}=0,\"לא בפיגור\",IF(G{i}<=30,\"1-30\",IF(G{i}<=60,\"31-60\",IF(G{i}<=90,\"61-90\",\"מעל 90\"))))",
                       f"=F{i}*VLOOKUP(H{i},'מדיניות הפרשה'!$A:$B,2,FALSE)"]
            else:
                ext = [None, None, None]
            rows.append(r + ext)
        table(ws2, 1, ["מס' חשבונית", "לקוח", "תאריך חשבונית", "תנאי תשלום", "תאריך פירעון", "יתרה פתוחה (כולל מע\"מ)",
                       "ימי פיגור", "דלי גיול", "הפרשה לחובות מסופקים"], rows,
              [12, 20, 14, 10, 14, 18, 10, 12, 18], {2: DATE, 4: DATE, 5: NUM, 8: NUM},
              input_cols=() if solution else (6, 7, 8))
        ws3 = new_sheet(wb, "מדיניות הפרשה")
        table(ws3, 1, ["דלי גיול", "שיעור הפרשה"], [(r[1], r[2]) for r in rates], [14, 14], {1: PCT})
        ws4 = new_sheet(wb, "סיכום גיול")
        agg = {}
        for r in inv_rows:
            days = max(0, (cutoff - r[4]).days)
            b, rt = bucket(days)
            a = agg.setdefault(b, [0, 0])
            a[0] += r[5]
            a[1] += r[5] * rt
        srows = []
        for i, (_, b, _) in enumerate(rates, start=2):
            if solution:
                srows.append([b, f"=SUMIFS('חשבוניות פתוחות'!F:F,'חשבוניות פתוחות'!H:H,A{i})",
                              f"=B{i}/$B$7", f"=SUMIFS('חשבוניות פתוחות'!I:I,'חשבוניות פתוחות'!H:H,A{i})"])
            else:
                srows.append([b, None, None, None])
        table(ws4, 1, ["דלי", "סה\"כ יתרה", "% מהסך", "הפרשה"], srows, [14, 16, 10, 14], {1: NUM, 2: PCT, 3: NUM},
              input_cols=() if solution else (1, 2, 3))
        ws4["A7"] = "סה\"כ"
        ws4["B7"] = "=SUM(B2:B6)"
        ws4["D7"] = "=SUM(D2:D6)"
        for c in ("B7", "D7"):
            ws4[c].number_format = NUM
            ws4[c].font = Font(bold=True)
        if solution:
            over = {}
            for r in inv_rows:
                if (cutoff - r[4]).days > 90:
                    over[r[1]] = over.get(r[1], 0) + r[5]
            top3 = sorted(over.items(), key=lambda x: -x[1])[:3]
            ws4["A10"] = "3 הלקוחות עם הכי הרבה חוב מעל 90 יום (אליהם מתקשרים קודם):"
            ws4["A10"].font = Font(bold=True)
            for k, (nm, v) in enumerate(top3, start=11):
                ws4.cell(row=k, column=1, value=nm)
                ws4.cell(row=k, column=2, value=v).number_format = NUM
            ws4["A15"] = "פקודת יומן להפרשה: חובה הוצאות חובות מסופקים / זכות הפרשה לחובות מסופקים (מוצג בקיזוז מהלקוחות במאזן)."

        ws5 = new_sheet(wb, "מע\"מ יוני")
        ws5["A1"], ws5["A1"].font = "חלק א: עסקאות (מכירות) יוני 2026, סכומים לפני מע\"מ", TITLE_FONT
        srows = []
        for i, (d, amt, kind) in enumerate(sales, start=3):
            f = (f"=IF(C{i}=\"יצוא\",0,B{i}*0.18)") if solution else None
            srows.append([d, amt, kind, f])
        table(ws5, 2, ["תיאור", "סכום לפני מע\"מ", "סוג", "מע\"מ עסקאות"], srows, [34, 16, 30, 16], {1: NUM, 3: NUM},
              input_cols=() if solution else (3,))
        top = 2 + len(sales) + 3
        ws5.cell(row=top - 1, column=1, value="חלק ב: תשומות (קניות) יוני 2026, סכומים לפני מע\"מ").font = TITLE_FONT
        prows = []
        for i, (d, amt, ok) in enumerate(purchases, start=top + 1):
            f1 = f"=B{i}*0.18" if solution else None
            f2 = f"=IF(LEFT(C{i},2)=\"כן\",D{i},0)" if solution else None
            prows.append([d, amt, ok, f1, f2])
        table(ws5, top, ["תיאור", "סכום לפני מע\"מ", "מותר לקיזוז?", "מע\"מ בחשבונית", "מע\"מ תשומות לקיזוז"], prows,
              [34, 16, 30, 16, 18], {1: NUM, 3: NUM, 4: NUM}, input_cols=() if solution else (3, 4))
        ws5.freeze_panes = None
        end = top + len(purchases)
        out_vat = sum(0 if k == "יצוא" else a * VAT for _, a, k in sales)
        in_vat = sum(a * VAT for _, a, ok in purchases if ok.startswith("כן"))
        summary = [("סה\"כ מע\"מ עסקאות", f"=SUM(D3:D{2 + len(sales)})", out_vat),
                   ("סה\"כ מע\"מ תשומות לקיזוז", f"=SUM(E{top + 1}:E{end})", in_vat),
                   ("מע\"מ לתשלום (שלילי = החזר)", f"=B{end + 3}-B{end + 4}", out_vat - in_vat)]
        for k, (lbl, f, v) in enumerate(summary, start=end + 3):
            ws5.cell(row=k, column=1, value=lbl).font = Font(bold=True)
            c = ws5.cell(row=k, column=2, value=f if solution else None)
            c.number_format = NUM
            c.fill = SOL_FILL if solution else INPUT_FILL
        if solution:
            ws5.cell(row=end + 7, column=1,
                     value=f"בדיקה: מע\"מ עסקאות {out_vat:,.0f}, תשומות {in_vat:,.0f}, לתשלום {out_vat - in_vat:,.0f} ₪")
        wb.save((SOL if solution else ROOT) / ("L05_receivables_aging_vat" + ("_solution" if solution else "") + ".xlsx"))


# ---------------------------------------------------------------------------
# שלב 6: התאמת בנק
# ---------------------------------------------------------------------------
def build_05():
    opening = 214300
    # ספרי החברה (כרטסת בנק בהנהלת חשבונות): (תאריך, אסמכתא, תיאור, חובה=כניסה, זכות=יציאה)
    books = [
        ("01/06/2026", "יתרת פתיחה", "יתרת פתיחה", opening, 0),
        ("02/06/2026", "קבלה 3301", "גבייה מרשת מזון השרון", 47200, 0),
        ("03/06/2026", "העברה 881", "תשלום לספק פולים עלית", 0, 38350),
        ("05/06/2026", "צ'ק 1044", "תשלום לספק כוסות", 0, 12600),
        ("09/06/2026", "העברה 882", "משכורות מאי", 0, 96400),
        ("12/06/2026", "קבלה 3302", "גבייה מסופר גליל", 31860, 0),
        ("15/06/2026", "העברה 883", "תשלום מע\"מ מאי", 0, 27940),
        ("16/06/2026", "צ'ק 1045", "שכירות יוני", 0, 25960),
        ("18/06/2026", "קבלה 3303", "גבייה ממלון הים", 18880, 0),
        ("22/06/2026", "העברה 884", "תשלום ביטוח לאומי וניכויים", 0, 21730),
        ("24/06/2026", "צ'ק 1046", "תשלום לספק עוגיות", 0, 4350),   # טעות: בפועל 3,450
        ("28/06/2026", "צ'ק 1047", "תשלום לרואה חשבון", 0, 8260),    # לא נפרע עדיין
        ("29/06/2026", "צ'ק 1048", "תשלום לספק מכסים", 0, 5900),     # לא נפרע עדיין
        ("30/06/2026", "קבלה 3304", "הפקדת צ'קים מקייטרינג זהב", 22420, 0),  # הפקדה בדרך
    ]
    # דף הבנק: (תאריך, תיאור, זכות=הפקדה לחשבון, חובה=משיכה)
    bank = [
        ("01/06/2026", "יתרת פתיחה", opening, 0),
        ("02/06/2026", "העברה נכנסת רשת מזון השרון", 47200, 0),
        ("03/06/2026", "העברה יוצאת עלית", 0, 38350),
        ("07/06/2026", "צ'ק 1044", 0, 12600),
        ("09/06/2026", "משכורות", 0, 96400),
        ("10/06/2026", "הוראת קבע ליסינג רכב", 0, 4720),           # לא נרשם בספרים
        ("12/06/2026", "העברה נכנסת סופר גליל", 31860, 0),
        ("15/06/2026", "מע\"מ", 0, 27940),
        ("17/06/2026", "צ'ק 1045", 0, 25960),
        ("18/06/2026", "העברה נכנסת מלון הים", 18880, 0),
        ("22/06/2026", "ביטוח לאומי", 0, 21730),
        ("26/06/2026", "צ'ק 1046", 0, 3450),
        ("30/06/2026", "עמלות ניהול חשבון", 0, 385),                # לא נרשם בספרים
        ("30/06/2026", "ריבית זכות", 142, 0),                       # לא נרשם בספרים
    ]
    book_bal = sum(r[3] - r[4] for r in books)
    bank_bal = sum(r[2] - r[3] for r in bank)
    adj_book = book_bal - 4720 - 385 + 142 + (4350 - 3450)
    adj_bank = bank_bal + 22420 - 8260 - 5900
    assert adj_book == adj_bank, (adj_book, adj_bank)
    for solution in (False, True):
        wb = Workbook()
        ws = new_sheet(wb, "הוראות", first=True)
        instructions(ws, "שלב 6 | התאמת בנק | אספקת בוקר בע\"מ | יוני 2026", [
            "הסיפור: היתרה בבנק לא שווה ליתרה בהנהלת החשבונות. זה נורמלי! התפקיד שלך: להסביר כל שקל של הפרש.",
            LEGEND,
            "",
            "שלב 1: חשבי את היתרה לסוף יוני בכל גיליון (ספרים ובנק). רשמי אותן בגיליון \"התאמה\".",
            "שלב 2: עברי שורה שורה. כל תנועה שמופיעה בשני הצדדים באותו סכום, סמני V בעמודת \"הותאם\".",
            "שלב 3: מה שנשאר בלי V הוא ההפרש. לכל שורה כזו שאלי: מי צודק? הבנק או אנחנו?",
            "שלב 4: בנק יודע ואנחנו לא (עמלות, ריבית, הוראת קבע) = מתקנים אצלנו בפקודת יומן.",
            "שלב 5: אנחנו יודעים והבנק עוד לא (צ'ק שלא נפרע, הפקדה בדרך) = לא מתקנים. רק מסבירים בהתאמה.",
            "שלב 6: טעות רישום אצלנו (סכום לא נכון) = מתקנים אצלנו.",
            "שלב 7: בסוף, יתרת ספרים מתואמת חייבת להיות שווה ליתרת בנק מתואמת. הפרש 0 = הצלחה.",
            "",
            "רמז: יש 7 הפרשים. חפשי: הוראת קבע, עמלה, ריבית, טעות הקלדה, 2 צ'קים, הפקדה אחת.",
        ])
        ws2 = new_sheet(wb, "ספרים (כרטסת בנק)")
        rows = [list(r) + [("V" if solution and r[1] not in ("צ'ק 1046", "צ'ק 1047", "צ'ק 1048", "קבלה 3304") else None)] for r in books]
        last = table(ws2, 1, ["תאריך", "אסמכתא", "תיאור", "חובה (נכנס)", "זכות (יצא)", "הותאם?"], rows,
                     [12, 14, 30, 14, 14, 10], {3: NUM, 4: NUM}, input_cols=() if solution else (5,))
        ws2.cell(row=last + 1, column=3, value="יתרה לסוף יוני").font = Font(bold=True)
        c = ws2.cell(row=last + 1, column=4, value=f"=SUM(D2:D{last})-SUM(E2:E{last})" if solution else None)
        c.number_format, c.fill = NUM, (SOL_FILL if solution else INPUT_FILL)
        ws3 = new_sheet(wb, "דף בנק")
        unmatched_bank = ("הוראת קבע ליסינג רכב", "עמלות ניהול חשבון", "ריבית זכות", "צ'ק 1046")
        rows = [list(r) + [("V" if solution and r[1] not in unmatched_bank else None)] for r in bank]
        last = table(ws3, 1, ["תאריך", "תיאור", "זכות (הפקדה)", "חובה (משיכה)", "הותאם?"], rows,
                     [12, 30, 14, 14, 10], {2: NUM, 3: NUM}, input_cols=() if solution else (4,))
        ws3.cell(row=last + 1, column=2, value="יתרה לסוף יוני").font = Font(bold=True)
        c = ws3.cell(row=last + 1, column=3, value=f"=SUM(C2:C{last})-SUM(D2:D{last})" if solution else None)
        c.number_format, c.fill = NUM, (SOL_FILL if solution else INPUT_FILL)

        ws4 = new_sheet(wb, "התאמה")
        ws4["A1"], ws4["A1"].font = "דוח התאמת בנק ליום 30.6.2026", TITLE_FONT
        left = [("יתרה לפי הספרים", book_bal), ("פחות: הוראת קבע ליסינג שלא נרשמה", -4720),
                ("פחות: עמלות בנק שלא נרשמו", -385), ("ועוד: ריבית זכות שלא נרשמה", 142),
                ("ועוד: תיקון טעות בצ'ק 1046 (נרשם 4,350 במקום 3,450)", 900),
                ("יתרת ספרים מתואמת", adj_book)]
        right = [("יתרה לפי דף הבנק", bank_bal), ("ועוד: הפקדה בדרך (קבלה 3304)", 22420),
                 ("פחות: צ'ק 1047 שטרם נפרע", -8260), ("פחות: צ'ק 1048 שטרם נפרע", -5900),
                 ("", None), ("יתרת בנק מתואמת", adj_bank)]
        ws4["A3"], ws4["C3"] = "צד הספרים (מה מתקנים אצלנו)", "צד הבנק (רק הסבר, לא מתקנים)"
        ws4["A3"].font = ws4["C3"].font = Font(bold=True)
        for i, ((l1, v1), (l2, v2)) in enumerate(zip(left, right), start=4):
            ws4.cell(row=i, column=1, value=l1 if solution or i in (4, 9) else None)
            c = ws4.cell(row=i, column=2, value=v1 if solution else None)
            c.number_format, c.fill = NUM, (SOL_FILL if solution else INPUT_FILL)
            ws4.cell(row=i, column=3, value=l2 if solution or i in (4, 9) else None)
            if v2 is not None:
                c = ws4.cell(row=i, column=4, value=v2 if solution else None)
                c.number_format, c.fill = NUM, (SOL_FILL if solution else INPUT_FILL)
        ws4["A11"] = "הפרש (חייב 0)"
        ws4["A11"].font = Font(bold=True)
        ws4["B11"] = "=B9-D9"
        ws4["B11"].number_format = NUM
        for col, w in zip("ABCD", (52, 14, 40, 14)):
            ws4.column_dimensions[col].width = w
        if solution:
            ws4["A14"] = "פקודות יומן לתיקון (רק צד הספרים):"
            ws4["A14"].font = Font(bold=True)
            je = ["1. חובה הוצאות ליסינג רכב 4,720 / זכות בנק 4,720",
                  "2. חובה עמלות בנק (הוצאות מימון) 385 / זכות בנק 385",
                  "3. חובה בנק 142 / זכות הכנסות מימון 142",
                  "4. חובה בנק 900 / זכות ספקים (ספק עוגיות) 900  (ביטול העודף שנרשם)",
                  "שימי לב: בצ'ק 1046 שילמנו לספק 3,450 אבל רשמנו 4,350. לכן גם כרטסת הספק שגויה ומתקנים אותה."]
            for k, t in enumerate(je, start=15):
                ws4.cell(row=k, column=1, value=t)
        wb.save((SOL if solution else ROOT) / ("L06_bank_reconciliation" + ("_solution" if solution else "") + ".xlsx"))


# ---------------------------------------------------------------------------
# שלב 7: סגירת חודש
# ---------------------------------------------------------------------------
def build_06():
    assets = [
        ("RK-001", "מחשבים ניידים (10 יח')", "ציוד מחשוב", dt.date(2024, 3, 1), 72000, 0.33),
        ("RK-002", "שרת מקומי", "ציוד מחשוב", dt.date(2025, 1, 1), 45000, 0.33),
        ("RK-003", "ריהוט משרדי", "ריהוט", dt.date(2023, 7, 1), 38000, 0.07),
        ("RK-004", "מכונת אריזה", "מכונות", dt.date(2024, 9, 1), 186000, 0.10),
        ("RK-005", "מלגזה", "כלי רכב ושינוע", dt.date(2025, 6, 1), 124000, 0.15),
        ("RK-006", "שיפוץ מחסן", "שיפורים במושכר", dt.date(2025, 11, 1), 96000, 0.10),
        ("RK-007", "מקררי תצוגה (4)", "מכונות", dt.date(2026, 2, 1), 58000, 0.10),
        ("RK-008", "מדפסת משרדית", "ציוד מחשוב", dt.date(2026, 5, 1), 6400, 0.33),
    ]
    month_end = dt.date(2026, 6, 30)
    employees = [("דנה כהן", 18000, 9.5), ("יוסי לוי", 14000, 22), ("מיכל אברהם", 22000, 4), ("עומר חדד", 11500, 15),
                 ("נועה פרץ", 16500, 31), ("אבי מזרחי", 26000, 12)]
    checklist = [
        ("WD1", "וידוא שכל חשבוניות הספקים של החודש נקלטו", "מנהלת חשבונות ספקים"),
        ("WD1", "הפקת חשבוניות ללקוחות על כל מה שסופק החודש", "מנהלת חשבונות לקוחות"),
        ("WD1", "קליטת דפי בנק וכרטיסי אשראי", "מנהלת חשבונות"),
        ("WD2", "התאמות בנק לכל החשבונות", "עוזרת חשבת"),
        ("WD2", "קליטת פקודת שכר מהמשכורות", "חשבת שכר"),
        ("WD2", "רישום פחת חודשי", "חשבת"),
        ("WD2", "הפחתת הוצאות מראש (ביטוח, מנויים)", "חשבת"),
        ("WD3", "הפרשות להוצאות לשלם (חשמל, ייעוץ שלא הגיעה חשבונית)", "חשבת"),
        ("WD3", "הפרשה לחופשה ולהבראה", "חשבת"),
        ("WD3", "הכנסות מראש / הכנסות לקבל", "חשבת"),
        ("WD3", "ספירת מלאי או עדכון עלות המכר", "מנהל מחסן + חשבת"),
        ("WD3", "התאמת כרטסות ספקים מרכזיים", "עוזרת חשבת"),
        ("WD4", "התאמת מע\"מ: כרטסת מול דוח", "חשבת"),
        ("WD4", "בדיקת סבירות: השוואה לחודש קודם ולתקציב", "חשבת"),
        ("WD4", "טיפול ביתרות חריגות (יתרת חובה לספק, זכות ללקוח)", "עוזרת חשבת"),
        ("WD5", "הפקת דוח רווח והפסד ומאזן חודשי", "חשבת"),
        ("WD5", "כתיבת הסברים לסטיות מהותיות", "חשבת"),
        ("WD5", "הצגה למנהל הכספים וקבלת אישור", "חשבת + סמנכ\"ל כספים"),
        ("WD5", "נעילת התקופה במערכת (שלא ירשמו לחודש סגור)", "חשבת"),
    ]

    def months_between(a, b):
        return (b.year - a.year) * 12 + (b.month - a.month) + 1

    for solution in (False, True):
        wb = Workbook()
        ws = new_sheet(wb, "הוראות", first=True)
        instructions(ws, "שלב 7 | סגירת חודש | אספקת בוקר בע\"מ | יוני 2026", [
            "הסיפור: הגיע ה-1 ביולי. המנכ\"ל רוצה דוח רווח והפסד של יוני עד ה-5 ביולי. את מובילה את הסגירה.",
            LEGEND,
            "",
            "משימה 1: פחת. בגיליון \"רכוש קבוע\" חשבי פחת חודשי = עלות × שיעור שנתי / 12 (שיטת הקו הישר).",
            "   חשבי גם כמה חודשים הנכס בשימוש עד 30/06/2026 (כולל חודש הרכישה), פחת נצבר ועלות מופחתת.",
            "   אם הפחת הנצבר גדול מהעלות, הנכס כבר הופחת במלואו. אז פחת החודש = 0. (רמז: MIN)",
            "משימה 2: הוצאות מראש. ביוני 2026 שולם ביטוח שנתי 24,000 ₪ ל-12 חודשים (יולי 2026 עד יוני 2027)? לא!",
            "   קראי בעיון בגיליון \"הפרשות\": מתי שולם ולאיזו תקופה. חשבי כמה נרשם כהוצאה ביוני וכמה נשאר נכס.",
            "משימה 3: הוצאות לשלם. חשבון החשמל של יוני יגיע רק באוגוסט. הערכי לפי ממוצע 3 חודשים קודמים.",
            "משימה 4: הפרשה לחופשה. לכל עובד: ימי חופשה צבורים × שכר יומי (שכר חודשי / 22).",
            "משימה 5: כתבי את כל פקודות היומן בגיליון \"פקודות סגירה\".",
            "משימה 6: בגיליון \"צ'קליסט סגירה\" עברי על 19 המשימות. זה הכלי הכי חשוב שלך בעבודה אמיתית.",
            "",
            "טיפ ADHD: הצ'קליסט הוא החבר הכי טוב שלך. חשבת מעולה לא זוכרת הכל. היא בונה רשימות שזוכרות בשבילה.",
        ])
        ws2 = new_sheet(wb, "רכוש קבוע")
        rows = []
        for i, a in enumerate(assets, start=2):
            if solution:
                ext = [f"=E{i}*F{i}/12",
                       f"=(2026-YEAR(D{i}))*12+(6-MONTH(D{i}))+1",
                       f"=MIN(E{i},G{i}*H{i})",
                       f"=E{i}-I{i}",
                       f"=IF(I{i}>=E{i},IF((H{i}-1)*G{i}<E{i},E{i}-(H{i}-1)*G{i},0),G{i})"]
            else:
                ext = [None] * 5
            rows.append(list(a) + ext)
        last = table(ws2, 1, ["מס' נכס", "תיאור", "קבוצה", "תאריך רכישה", "עלות", "שיעור פחת שנתי",
                              "פחת חודשי", "חודשים בשימוש", "פחת נצבר", "עלות מופחתת", "פחת לרשום ביוני"], rows,
                     [9, 22, 16, 13, 12, 10, 12, 10, 12, 12, 14], {3: DATE, 4: NUM, 5: PCT, 6: NUM2, 8: NUM, 9: NUM, 10: NUM2},
                     input_cols=() if solution else (6, 7, 8, 9, 10))
        ws2.cell(row=last + 1, column=2, value="סה\"כ").font = Font(bold=True)
        for col in "EGIJK":
            c = ws2[f"{col}{last + 1}"]
            c.value = f"=SUM({col}2:{col}{last})"
            c.number_format = NUM
        ws3 = new_sheet(wb, "הפרשות")
        ws3["A1"], ws3["A1"].font = "נתונים להפרשות סוף חודש", TITLE_FONT
        info = [
            ("ביטוח שנתי: שולם ב-01/01/2026 סכום 24,000 ₪ לתקופה 01/01/2026 עד 31/12/2026", None),
            ("הוצאת ביטוח לחודש יוני", "=24000/12"),
            ("יתרת ביטוח מראש (נכס) ב-30/06/2026 (6 חודשים שנותרו)", "=24000/12*6"),
            ("", None),
            ("חשמל: חשבון מרץ", 3850), ("חשמל: חשבון אפריל", 4120), ("חשמל: חשבון מאי", 4460),
            ("הערכת חשמל יוני (ממוצע 3 חודשים) = הוצאות לשלם", "=ROUND(AVERAGE(B7:B9),0)"),
            ("", None),
            ("ייעוץ משפטי: עורך הדין עבד ביוני 12 שעות × 650 ₪. חשבונית עוד לא הגיעה", "=12*650"),
        ]
        for i, (lbl, v) in enumerate(info, start=3):
            ws3.cell(row=i, column=1, value=lbl)
            if v is not None:
                is_given = isinstance(v, (int, float))
                c = ws3.cell(row=i, column=2, value=v if (solution or is_given) else None)
                c.number_format = NUM
                if not is_given:
                    c.fill = SOL_FILL if solution else INPUT_FILL
        ws3["A15"] = "הפרשה לחופשה"
        ws3["A15"].font = Font(bold=True)
        erows = []
        for i, (nm, sal, days) in enumerate(employees, start=17):
            erows.append([nm, sal, days] + ([f"=B{i}/22", f"=C{i}*D{i}"] if solution else [None, None]))
        last = table(ws3, 16, ["עובד", "שכר חודשי", "ימי חופשה צבורים", "שכר יומי", "הפרשה"], erows,
                     [70, 12, 16, 12, 12], {1: NUM, 3: NUM2, 4: NUM}, input_cols=() if solution else (3, 4))
        ws3.freeze_panes = None
        ws3.cell(row=last + 1, column=1, value="סה\"כ הפרשה לחופשה (יתרה נדרשת)").font = Font(bold=True)
        c = ws3.cell(row=last + 1, column=5, value=f"=SUM(E17:E{last})")
        c.number_format = NUM
        ws3.cell(row=last + 2, column=1, value="יתרת ההפרשה בספרים לפני הסגירה")
        ws3.cell(row=last + 2, column=5, value=64000).number_format = NUM
        ws3.cell(row=last + 3, column=1, value="הפרשה לרשום ביוני = יתרה נדרשת פחות יתרה קיימת").font = Font(bold=True)
        c = ws3.cell(row=last + 3, column=5, value=f"=E{last + 1}-E{last + 2}" if solution else None)
        c.number_format, c.fill = NUM, (SOL_FILL if solution else INPUT_FILL)

        ws4 = new_sheet(wb, "פקודות סגירה")
        if solution:
            dep_total = 0.0
            for a in assets:
                monthly = a[4] * a[5] / 12
                m = months_between(a[3], month_end)
                if monthly * m <= a[4]:
                    dep_total += monthly
                elif monthly * (m - 1) < a[4]:
                    dep_total += a[4] - monthly * (m - 1)
            vac_req = sum(s / 22 * d for _, s, d in employees)
            je = [
                (1, "הוצאות פחת", round(dep_total), None), (1, "פחת נצבר", None, round(dep_total)),
                (2, "הוצאות ביטוח", 2000, None), (2, "ביטוח מראש", None, 2000),
                (3, "הוצאות חשמל", round((3850 + 4120 + 4460) / 3), None), (3, "הוצאות לשלם", None, round((3850 + 4120 + 4460) / 3)),
                (4, "הוצאות ייעוץ משפטי", 7800, None), (4, "הוצאות לשלם", None, 7800),
                (5, "הוצאות שכר (חופשה)", round(vac_req - 64000), None), (5, "הפרשה לחופשה", None, round(vac_req - 64000)),
            ]
        else:
            je = [(k, None, None, None) for k in (1, 1, 2, 2, 3, 3, 4, 4, 5, 5)]
        table(ws4, 1, ["פקודה", "חשבון", "חובה", "זכות"], je, [8, 26, 14, 14], {2: NUM, 3: NUM},
              input_cols=() if solution else (1, 2, 3))
        ws5 = new_sheet(wb, "צ'קליסט סגירה")
        rows = [[d, t, o, None, None] for d, t, o in checklist]
        table(ws5, 1, ["יום עבודה", "משימה", "אחראי", "בוצע? (V)", "הערות / חריגים"], rows, [10, 58, 24, 10, 30],
              input_cols=(3, 4))
        wb.save((SOL if solution else ROOT) / ("L07_month_end_close" + ("_solution" if solution else "") + ".xlsx"))


# ---------------------------------------------------------------------------
# שלב 8: תקציב מול ביצוע ותזרים
# ---------------------------------------------------------------------------
def build_07():
    lines = [
        ("הכנסות", "הכנסה", 1850000, 1712400),
        ("עלות המכר", "הוצאה", 1110000, 1068300),
        ("שכר ונלוות", "הוצאה", 318000, 331500),
        ("שכירות ואחזקה", "הוצאה", 108000, 108000),
        ("שיווק ופרסום", "הוצאה", 92000, 141800),
        ("הוצאות משרד", "הוצאה", 31000, 28700),
        ("שירותים מקצועיים", "הוצאה", 24000, 46500),
        ("פחת", "הוצאה", 62000, 62400),
        ("הוצאות מימון", "הוצאה", 20500, 23900),
    ]
    weeks = 13
    rnd = random.Random(11)
    for solution in (False, True):
        wb = Workbook()
        ws = new_sheet(wb, "הוראות", first=True)
        instructions(ws, "שלב 8 | תקציב מול ביצוע ותזרים מזומנים | אספקת בוקר בע\"מ", [
            "הסיפור: ישיבת הנהלה ביום חמישי. סמנכ\"ל הכספים מבקש: 1) הסבר לסטיות ברבעון 2. 2) תחזית מזומנים ל-13 שבועות.",
            LEGEND,
            "",
            "חלק א: תקציב מול ביצוע (רבעון 2, אפריל עד יוני 2026)",
            "שלב 1: סטייה בש\"ח = ביצוע פחות תקציב.",
            "שלב 2: סטייה באחוזים = סטייה / תקציב.",
            "שלב 3: חיובית או שלילית לעסק? בהכנסות: ביצוע גבוה = טוב. בהוצאות: ביצוע גבוה = רע. כתבי \"חיובית\" או \"שלילית\".",
            "שלב 4: מהותית? כלל אצלנו: סטייה מעל 10% וגם מעל 20,000 ₪. (רמז: AND ו-ABS)",
            "שלב 5: לכל סטייה מהותית כתבי הסבר בשפה פשוטה (השתמשי ברמזים בגיליון \"מידע מהמחלקות\").",
            "",
            "חלק ב: תזרים מזומנים ל-13 שבועות",
            "שלב 6: בגיליון \"תזרים 13 שבועות\" יתרת סגירה = יתרת פתיחה + כניסות - יציאות.",
            "שלב 7: יתרת הפתיחה של כל שבוע = יתרת הסגירה של השבוע הקודם.",
            "שלב 8: המדיניות: היתרה לא יורדת מתחת ל-150,000 ₪. באילו שבועות יש בעיה? למה? מה היית מציעה לעשות?",
            "",
            "למה זה חשוב? זה הרגע שבו חשבת הופכת משותפה עסקית ולא רק \"מי שרושמת\". כאן רואים את הערך שלך.",
        ])
        ws2 = new_sheet(wb, "תקציב מול ביצוע")
        rows = []
        for i, (nm, kind, b, a) in enumerate(lines, start=2):
            if solution:
                ext = [f"=D{i}-C{i}", f"=E{i}/C{i}",
                       f"=IF(B{i}=\"הכנסה\",IF(E{i}>=0,\"חיובית\",\"שלילית\"),IF(E{i}<=0,\"חיובית\",\"שלילית\"))",
                       f"=IF(AND(ABS(F{i})>10%,ABS(E{i})>20000),\"מהותית\",\"\")"]
            else:
                ext = [None] * 4
            rows.append([nm, kind, b, a] + ext + [None])
        last = table(ws2, 1, ["סעיף", "סוג", "תקציב רבעון 2", "ביצוע רבעון 2", "סטייה ₪", "סטייה %", "חיובית/שלילית",
                              "מהותית?", "הסבר להנהלה"], rows, [20, 8, 14, 14, 12, 10, 14, 10, 60],
                     {2: NUM, 3: NUM, 4: NUM, 5: PCT}, input_cols=() if solution else (4, 5, 6, 7, 8))
        ws2.cell(row=last + 2, column=1, value="רווח תפעולי ומימון (הכנסות פחות כל ההוצאות)").font = Font(bold=True)
        for col in "CD":
            c = ws2[f"{col}{last + 2}"]
            c.value = f"={col}2-SUM({col}3:{col}{last})"
            c.number_format = NUM
            c.font = Font(bold=True)
        if solution:
            expl = {
                "הכנסות": "לא מהותית לפי הכלל (7.4%, מתחת ל-10%), אבל שווה משפט: סופר גליל דחה הזמנה של כ-120 אלף ₪ ליולי. נדחה, לא אבוד.",
                "שיווק ופרסום": "חריגה של 54%. קמפיין השקה של קו הקרואסונים שלא היה בתקציב. לבדוק החזר השקעה ברבעון 3.",
                "שירותים מקצועיים": "כמעט פי 2 מהתקציב. ייעוץ משפטי חד פעמי בתביעת ספק. לא צפוי לחזור.",
            }
            for i, (nm, *_rest) in enumerate(lines, start=2):
                if nm in expl:
                    ws2.cell(row=i, column=9, value=expl[nm]).fill = SOL_FILL
        ws3 = new_sheet(wb, "מידע מהמחלקות")
        instructions(ws3, "מה סיפרו לך המנהלים (רמזים להסברים)", [
            "מנהל המכירות: \"סופר גליל דחה הזמנה גדולה של כ-120 אלף ₪ ליולי. הם יקנו, רק מאוחר יותר.\"",
            "מנהלת השיווק: \"עשינו קמפיין השקה לקרואסונים. זה לא היה בתקציב המקורי, המנכ\"ל אישר בעל פה.\"",
            "היועץ המשפטי: \"התביעה של הספק לשעבר דרשה עבודה רבה ברבעון. זה חד פעמי.\"",
            "משאבי אנוש: \"גייסנו נהג נוסף במאי במקום לשלם שעות נוספות.\" (האם זו סטייה מהותית? בדקי לפי הכלל!)",
        ])
        ws4 = new_sheet(wb, "תזרים 13 שבועות")
        start = dt.date(2026, 7, 5)
        rows = []
        opening = 250000
        for w in range(weeks):
            d = start + dt.timedelta(days=7 * w)
            coll = [142000, 118000, 96000, 131000, 88000, 125000, 99000, 140000, 91000, 117000, 136000, 104000, 128000][w] + 60000
            payroll = 118000 if w in (0, 4, 8, 12) else 0
            suppliers = [96000, 104000, 88000, 97000, 101000, 93000, 110000, 99000, 86000, 102000, 95000, 108000, 90000][w]
            vat = 61000 if w in (1, 5, 10) else 0
            rent = 36000 if w in (0, 4, 8) else 0
            loan = 45000 if w == 6 else 0
            tax_adv = 22000 if w in (2, 6, 11) else 0
            other = 0 if w != 5 else 140000  # תשלום חד פעמי: מקדמה על מכונה
            r = 2 + w
            if solution:
                op = f"={opening}" if w == 0 else f"=K{r - 1}"
                close = f"=B{r}+C{r}-SUM(D{r}:I{r})"
                flag = f"=IF(K{r}<150000,\"מתחת למינימום!\",\"\")"
            else:
                op = opening if w == 0 else None
                close = flag = None
            rows.append([f"שבוע {w + 1} ({d:%d/%m})", op, coll, payroll, suppliers, vat, rent, loan, tax_adv, other, close, flag])
        # שורה J = אחר, נכלל בסכום: נתקן טווח
        for rr in rows:
            if solution and isinstance(rr[10], str):
                rr[10] = rr[10].replace("SUM(D", "SUM(D").replace(":I", ":J")
        table(ws4, 1, ["שבוע", "יתרת פתיחה", "גבייה מלקוחות", "שכר", "ספקים", "מע\"מ", "שכירות", "החזר הלוואה",
                       "מקדמות מס", "אחר (מקדמה על מכונה)", "יתרת סגירה", "התראה"], rows,
              [16, 14, 14, 11, 11, 10, 10, 12, 12, 16, 14, 16], {i: NUM for i in range(1, 11)},
              input_cols=() if solution else (1, 10, 11))
        if solution:
            bal = opening
            low = []
            for w, rr in enumerate(rows):
                bal = bal + rr[2] - sum(rr[3:10])
                if bal < 150000:
                    low.append((w + 1, bal))
            msg = "; ".join(f"שבוע {w}: {b:,.0f} ₪" for w, b in low) or "אין"
            ws4.cell(row=17, column=1, value="שבועות מתחת למינימום: " + msg).font = Font(bold=True)
            ws4.cell(row=18, column=1, value="הצעות: שבועות 6-7 בגלל מקדמה חד פעמית על מכונה, שבוע 9 בגלל שכר ושכירות באותו שבוע. לדחות את המקדמה או לפרוס אותה, לבקש מסופר גליל לשלם מראש, "
                                            "לפרוס תשלום לספק גדול, או להשתמש במסגרת אשראי בבנק לתקופה קצרה.")
        wb.save((SOL if solution else ROOT) / ("L08_budget_vs_actual_cashflow" + ("_solution" if solution else "") + ".xlsx"))


# ---------------------------------------------------------------------------
# שלב 9: איחוד דוחות (קבוצה עם חברה בת בחו"ל)
# ---------------------------------------------------------------------------
def build_08():
    closing, average, hist = 3.70, 3.65, 3.60
    parent = [
        ("מזומנים", "נכס", 1240000), ("לקוחות", "נכס", 2310000), ("חייבים חברה קשורה (Gilboa Inc)", "נכס", 185000),
        ("השקעה בחברה בת", "נכס", 360000), ("רכוש קבוע נטו", "נכס", 1650000),
        ("ספקים", "התחייבות", -1420000), ("הלוואות", "התחייבות", -900000),
        ("הון מניות", "הון", -500000), ("עודפים לתחילת שנה", "הון", None),
        ("הכנסות", "רוו\"ה", -9800000), ("הכנסות מחברה בת", "רוו\"ה", -438000),
        ("עלות המכר", "רוו\"ה", 5900000), ("הוצאות תפעול", "רוו\"ה", 3120000), ("הוצאות מימון", "רוו\"ה", 64000),
    ]
    plug = -sum(v for _, _, v in parent if v is not None)
    parent = [(a, b, plug if v is None else v) for a, b, v in parent]
    sub = [  # דולר
        ("מזומנים", "נכס", 210000, "סגירה"), ("לקוחות", "נכס", 340000, "סגירה"),
        ("רכוש קבוע נטו", "נכס", 95000, "סגירה"),
        ("ספקים", "התחייבות", -180000, "סגירה"), ("זכאים חברה אם", "התחייבות", -50000, "סגירה"),
        ("הון מניות", "הון", -100000, "היסטורי"), ("עודפים לתחילת שנה", "הון", -155000, "היסטורי*"),
        ("הכנסות", "רוו\"ה", -1400000, "ממוצע"), ("עלות המכר", "רוו\"ה", 690000, "ממוצע"),
        ("שירותים מחברה אם", "רוו\"ה", 120000, "ממוצע"), ("הוצאות תפעול", "רוו\"ה", 430000, "ממוצע"),
    ]
    assert sum(v for _, _, v, _ in sub) == 0
    re_open_ils = 543000  # עודפי פתיחה של הבת בשקלים (לפי שערים היסטוריים, נתון)
    rate_map = {"סגירה": closing, "ממוצע": average, "היסטורי": hist}

    def ils(v, r):
        if r == "היסטורי*":
            return -re_open_ils if v < 0 else re_open_ils
        return round(v * rate_map[r])

    sub_ils = [(a, b, v, r, ils(v, r)) for a, b, v, r in sub]
    trans_diff = -sum(x[4] for x in sub_ils)  # הפרש תרגום לקרן בהון
    for solution in (False, True):
        wb = Workbook()
        ws = new_sheet(wb, "הוראות", first=True)
        instructions(ws, "שלב 9 | איחוד דוחות כספיים | קבוצת גלבוע טכנולוגיות | 31.12.2026", [
            "הסיפור: לחברה הישראלית (חברה אם) יש חברה בת בארה\"ב שמחזיקה בה 100%. צריך דוח אחד לכל הקבוצה.",
            "זו משימה של חשבת בכירה בחברה גדולה או ציבורית. אם זה נראה מפחיד, זה בסדר. עושים צעד אחד בכל פעם.",
            LEGEND,
            "",
            "שלב 1: תרגום. החברה הבת מדווחת בדולרים. תרגמי לשקלים בגיליון \"בת - תרגום\":",
            "   נכסים והתחייבויות = שער סגירה 3.70. הכנסות והוצאות = שער ממוצע 3.65. הון מניות = שער היסטורי 3.60.",
            "   עודפי פתיחה בשקלים נתונים (543,000 ₪). ההפרש שנוצר בתרגום עובר ל\"קרן הפרשי תרגום\" בהון.",
            "שלב 2: חיבור. בגיליון \"איחוד\" חברי שורה שורה: חברה אם + בת בשקלים.",
            "שלב 3: ביטולים (הכי חשוב!). הקבוצה לא יכולה להרוויח מעצמה או להיות חייבת לעצמה:",
            "   א. לבטל חוב בין החברות: חייבים חברה קשורה (185,000) מול זכאים חברה אם (185,000).",
            "   ב. לבטל הכנסות והוצאות בין החברות: הכנסות מחברה בת (438,000) מול שירותים מחברה אם (438,000).",
            "   ג. לבטל את ההשקעה בבת (360,000) מול הון המניות של הבת (360,000).",
            "שלב 4: מאוחד = סכום פלוס ביטולים. בדקי שסך המאוחד עדיין 0 (מאוזן).",
            "שלב 5: חשבי את הרווח הנקי המאוחד. (הוא צריך להיות רווח האם + רווח הבת, בלי הכנסות פנימיות.)",
        ])
        ws2 = new_sheet(wb, "חברה אם (₪)")
        last_p = table(ws2, 1, ["סעיף", "סוג", "יתרה ₪ (חובה + / זכות -)"], [list(x) for x in parent], [34, 10, 22], {2: NUM})
        ws2.cell(row=last_p + 1, column=1, value="בדיקה (0)")
        ws2.cell(row=last_p + 1, column=3, value=f"=SUM(C2:C{last_p})").number_format = NUM
        ws3 = new_sheet(wb, "בת - תרגום")
        rows = []
        for i, (a, b, v, r, s) in enumerate(sub_ils, start=2):
            f = s if solution else None
            rows.append([a, b, v, r, (rate_map.get(r, "נתון")), f])
        last_s = table(ws3, 1, ["סעיף", "סוג", "יתרה $", "סוג שער", "שער", "יתרה ₪"], rows, [26, 10, 14, 12, 8, 16],
                       {2: NUM, 5: NUM}, input_cols=() if solution else (5,))
        ws3.cell(row=last_s + 1, column=1, value="קרן הפרשי תרגום (משלים לאפס)")
        c = ws3.cell(row=last_s + 1, column=6, value=trans_diff if solution else None)
        c.number_format, c.fill = NUM, (SOL_FILL if solution else INPUT_FILL)
        ws3.cell(row=last_s + 2, column=1, value="בדיקה (0)")
        ws3.cell(row=last_s + 2, column=6, value=f"=SUM(F2:F{last_s + 1})").number_format = NUM

        ws4 = new_sheet(wb, "איחוד")
        keys = ["מזומנים", "לקוחות", "חייבים חברה קשורה (Gilboa Inc)", "השקעה בחברה בת", "רכוש קבוע נטו",
                "ספקים", "זכאים חברה אם", "הלוואות", "הון מניות", "עודפים לתחילת שנה", "קרן הפרשי תרגום",
                "הכנסות", "הכנסות מחברה בת", "שירותים מחברה אם", "עלות המכר", "הוצאות תפעול", "הוצאות מימון"]
        pmap = {a: v for a, _, v in parent}
        smap = {a: s for a, _, _, _, s in sub_ils}
        smap["קרן הפרשי תרגום"] = trans_diff
        elim = {"חייבים חברה קשורה (Gilboa Inc)": -185000, "זכאים חברה אם": 185000,
                "הכנסות מחברה בת": 438000, "שירותים מחברה אם": -438000,
                "השקעה בחברה בת": -360000, "הון מניות": 360000}
        rows = []
        for i, k in enumerate(keys, start=2):
            p = pmap.get(k, 0)
            s = smap.get(k, 0)
            e = elim.get(k, 0)
            rows.append([k, p, (s if solution else None), (f"=B{i}+C{i}" if solution else None),
                         (e if solution else None), (f"=D{i}+E{i}" if solution else None)])
        last = table(ws4, 1, ["סעיף", "חברה אם ₪", "חברה בת ₪", "סכום", "ביטולים", "מאוחד"], rows,
                     [34, 14, 14, 14, 14, 14], {1: NUM, 2: NUM, 3: NUM, 4: NUM, 5: NUM},
                     input_cols=() if solution else (2, 3, 4, 5))
        ws4.cell(row=last + 1, column=1, value="בדיקה: סך הכל (0)").font = Font(bold=True)
        for col in "BCDEF":
            ws4[f"{col}{last + 1}"] = f"=SUM({col}2:{col}{last})"
            ws4[f"{col}{last + 1}"].number_format = NUM
        ws4.cell(row=last + 3, column=1, value="רווח נקי מאוחד (הכנסות פחות הוצאות)").font = Font(bold=True)
        c = ws4.cell(row=last + 3, column=6, value=f"=-SUM(F13:F{last})" if solution else None)
        c.number_format, c.fill = NUM, (SOL_FILL if solution else INPUT_FILL)
        if solution:
            p_net = -sum(v for a, t, v in parent if t == "רוו\"ה")
            s_net = -sum(s for a, t, v, r, s in sub_ils if t == "רוו\"ה")
            ws4.cell(row=last + 5, column=1,
                     value=f"בדיקה: רווח אם {p_net:,.0f} + רווח בת {s_net:,.0f} = {p_net + s_net:,.0f} ₪. "
                           f"הביטול של הכנסות/הוצאות פנימיות לא משנה את הרווח, רק מקטין הכנסות והוצאות באותו סכום.")
        wb.save((SOL if solution else ROOT) / ("L09_consolidation" + ("_solution" if solution else "") + ".xlsx"))


# ---------------------------------------------------------------------------
# שלב 11: פרויקט גמר. ייצוא ספר ראשי עם חריגות מוסתרות
# ---------------------------------------------------------------------------
def build_09():
    rnd = random.Random(2026)
    suppliers = [("S01", "עלית חומרי גלם"), ("S02", "כוסות ישראל"), ("S03", "הובלות צפון"), ("S04", "חשמל ישראל"),
                 ("S05", "משרד עו\"ד כהן"), ("S06", "גוגל פרסום"), ("S07", "נכסי השרון (שכירות)"), ("S08", "מחשבים פלוס")]
    acc = {"4000": "הכנסות", "5000": "עלות המכר", "6100": "שכר", "6200": "שכירות", "6300": "חשמל",
           "6400": "שיווק", "6500": "משרד", "6600": "שירותים מקצועיים", "6700": "הובלות",
           "1010": "בנק", "1100": "לקוחות", "2000": "ספקים", "2100": "מע\"מ", "2200": "מוסדות שכר"}
    lines = []
    ref = 90001

    def add(date, code, desc, d, c, dept, partner=""):
        lines.append([ref, date, code, acc[code], desc, d, c, dept, partner])

    days = [dt.date(2026, 7, d) for d in range(1, 32)]
    workdays = [d for d in days if d.weekday() not in (4, 5)]  # שישי ושבת
    # מכירות
    for d in workdays:
        amt = rnd.randint(18, 60) * 1000
        add(d, "1100", "חשבונית מכירה", round(amt * 1.18), 0, "מכירות", "לקוחות שונים")
        add(d, "4000", "חשבונית מכירה", 0, amt, "מכירות", "לקוחות שונים")
        add(d, "2100", "מע\"מ עסקאות", 0, round(amt * 0.18), "מכירות", "")
        ref += 1
    # קניות וספקים
    for d in workdays[::2]:
        s = rnd.choice([suppliers[0], suppliers[1], suppliers[2]])
        amt = rnd.randint(8, 30) * 1000
        code = "5000" if s[0] in ("S01", "S02") else "6700"
        add(d, code, f"חשבונית ספק {s[1]}", amt, 0, "תפעול", s[0])
        add(d, "2100", "מע\"מ תשומות", round(amt * 0.18), 0, "תפעול", s[0])
        add(d, "2000", f"חשבונית ספק {s[1]}", 0, round(amt * 1.18), "תפעול", s[0])
        ref += 1
    fixed = [(dt.date(2026, 7, 1), "6200", "S07", 36000, "הנהלה"), (dt.date(2026, 7, 14), "6300", "S04", 4380, "הנהלה"),
             (dt.date(2026, 7, 20), "6400", "S06", 15600, "שיווק"), (dt.date(2026, 7, 23), "6600", "S05", 7800, "הנהלה"),
             (dt.date(2026, 7, 27), "6500", "S08", 5200, "הנהלה")]
    for d, code, sup, amt, dept in fixed:
        add(d, code, f"חשבונית {dict(suppliers)[sup]}", amt, 0, dept, sup)
        add(d, "2100", "מע\"מ תשומות", round(amt * 0.18), 0, dept, sup)
        add(d, "2000", f"חשבונית {dict(suppliers)[sup]}", 0, round(amt * 1.18), dept, sup)
        ref += 1
    # שכר
    for dept, amt in (("הנהלה", 88000), ("מכירות", 64000), ("תפעול", 97000), ("שיווק", 31000)):
        add(dt.date(2026, 7, 31), "6100", f"שכר יולי {dept}", amt, 0, dept, "")
        add(dt.date(2026, 7, 31), "2200", f"שכר יולי {dept}", 0, amt, dept, "")
        ref += 1
    # חריגות מוסתרות
    anomalies = []
    # 1. תשלום כפול לספק
    for _ in range(2):
        add(dt.date(2026, 7, 16), "2000", "תשלום לספק הובלות צפון", 23600, 0, "תפעול", "S03")
        add(dt.date(2026, 7, 16), "1010", "תשלום לספק הובלות צפון", 0, 23600, "תפעול", "S03")
        ref += 1
    anomalies.append(f"1. תשלום כפול: פקודות {ref - 2} ו-{ref - 1}, ספק הובלות צפון, 23,600 ₪ באותו יום.")
    # 2. פקודה לא מאוזנת
    add(dt.date(2026, 7, 21), "6500", "ציוד משרדי", 3400, 0, "הנהלה", "S08")
    add(dt.date(2026, 7, 21), "2000", "ציוד משרדי", 0, 3040, "הנהלה", "S08")
    anomalies.append(f"2. פקודה לא מאוזנת: פקודה {ref}, חובה 3,400 מול זכות 3,040 (היפוך ספרות).")
    ref += 1
    # 3. תאריך בשבת וסכום עגול לספק לא מוכר
    sat = [d for d in days if d.weekday() == 5][2]
    add(sat, "6600", "ייעוץ", 50000, 0, "הנהלה", "S99")
    add(sat, "1010", "ייעוץ", 0, 50000, "הנהלה", "S99")
    anomalies.append(f"3. תשלום חשוד: פקודה {ref}, 50,000 ₪ עגול, ביום שבת {sat:%d/%m}, לספק S99 שלא קיים ברשימת הספקים ובלי חשבונית ומע\"מ.")
    ref += 1
    # 4. הוצאה שנרשמה לתקופה לא נכונה (תאריך אוגוסט)
    add(dt.date(2026, 8, 3), "6200", "שכירות אוגוסט", 36000, 0, "הנהלה", "S07")
    add(dt.date(2026, 8, 3), "2000", "שכירות אוגוסט", 0, 36000, "הנהלה", "S07")
    anomalies.append(f"4. תקופה שגויה: פקודה {ref} עם תאריך 03/08/2026 (שכירות אוגוסט) נמצאת בייצוא של יולי. גם חסר מע\"מ.")
    ref += 1
    # 5. הכנסה במחלקה לא נכונה + סכום חריג פי 10
    add(dt.date(2026, 7, 29), "1100", "חשבונית מכירה", 531000, 0, "הנהלה", "לקוחות שונים")
    add(dt.date(2026, 7, 29), "4000", "חשבונית מכירה", 0, 450000, "הנהלה", "לקוחות שונים")
    add(dt.date(2026, 7, 29), "2100", "מע\"מ עסקאות", 0, 81000, "הנהלה", "")
    anomalies.append(f"5. סכום חריג: פקודה {ref}, מכירה של 450,000 ₪ ביום אחד (פי 10 מהרגיל) ורשומה במחלקת הנהלה ולא מכירות. אולי אפס מיותר (45,000).")
    ref += 1
    rnd.shuffle(lines)
    lines.sort(key=lambda r: (r[1], r[0]))
    for solution in (False, True):
        wb = Workbook()
        ws = new_sheet(wb, "הוראות", first=True)
        instructions(ws, "שלב 11 | פרויקט גמר | דוח חודשי להנהלה מייצוא גולמי | יולי 2026", [
            "הסיפור: זה היום הראשון שלך כחשבת. קיבלת ייצוא של כל התנועות של יולי מהמערכת (כמו מפריוריטי או SAP).",
            "המנכ\"ל רוצה עד מחר: דוח רווח והפסד של יולי לפי מחלקות, ו\"האם יש משהו מוזר שאני צריך לדעת?\"",
            LEGEND,
            "",
            "שלב 1: בדקי שהקובץ מאוזן: סך חובה = סך זכות. אם לא, מצאי איזו פקודה לא מאוזנת (רמז: SUMIFS לפי אסמכתא).",
            "שלב 2: בני מאזן בוחן לפי חשבון (טבלת ציר: שורות = קוד חשבון, ערכים = סכום חובה וסכום זכות).",
            "שלב 3: בני דוח רווח והפסד של יולי לפי מחלקות (טבלת ציר: שורות = חשבון, עמודות = מחלקה).",
            "שלב 4: חפשי 5 חריגות שהוסתרו בקובץ. חשבי כמו מבקרת: כפילויות, ימים לא הגיוניים, סכומים עגולים,",
            "   ספקים לא מוכרים, תאריכים מחוץ לחודש, סכומים חריגים, פקודות לא מאוזנות.",
            "שלב 5: כתבי בגיליון \"דוח להנהלה\" עמוד אחד: 3 מספרים חשובים, 5 החריגות ומה מציעים לעשות.",
            "",
            "בונוס AI: העלי לכלי בינה מלאכותית (מאושר בחברה!) רק את מבנה הקובץ, בלי שמות ספקים אמיתיים, ובקשי ממנו",
            "   להציע בדיקות נוספות. השווי לרשימה שלך. מה הוא מצא שאת לא? מה את מצאת שהוא לא?",
        ])
        ws2 = new_sheet(wb, "תנועות יולי")
        table(ws2, 1, ["אסמכתא", "תאריך", "קוד חשבון", "שם חשבון", "תיאור", "חובה", "זכות", "מחלקה", "ספק/לקוח"],
              lines, [10, 12, 10, 16, 30, 12, 12, 10, 14], {1: DATE, 5: NUM, 6: NUM})
        ws2.auto_filter.ref = f"A1:I{len(lines) + 1}"
        ws3 = new_sheet(wb, "ספקים מאושרים")
        table(ws3, 1, ["קוד ספק", "שם ספק"], suppliers, [10, 24])
        ws4 = new_sheet(wb, "דוח להנהלה")
        ws4["A1"], ws4["A1"].font = "דוח חודשי להנהלה | יולי 2026", TITLE_FONT
        ws4.column_dimensions["A"].width = 120
        if solution:
            rev = sum(l[6] for l in lines if l[2] == "4000")
            exp = sum(l[5] for l in lines if l[2] in ("5000", "6100", "6200", "6300", "6400", "6500", "6600", "6700"))
            dr = sum(l[5] for l in lines)
            cr = sum(l[6] for l in lines)
            txt = [
                f"סך חובה {dr:,.0f} מול סך זכות {cr:,.0f}. הפרש {dr - cr:,.0f} ₪ = הפקודה הלא מאוזנת (חריגה 2).",
                f"הכנסות לפי הקובץ: {rev:,.0f} ₪. אם החשבונית של 450 אלף היא בעצם 45 אלף (חריגה 5), ההכנסות האמיתיות נמוכות ב-405,000 ₪.",
                f"הוצאות לפי הקובץ (כולל שכירות אוגוסט שלא שייכת ליולי): {exp:,.0f} ₪.",
                "",
                "5 החריגות:",
            ] + anomalies + [
                "",
                "המלצות: 1) לבקש החזר מהובלות צפון על התשלום הכפול. 2) לתקן את הפקודה הלא מאוזנת ולנעול במערכת רישום לא מאוזן.",
                "3) לעצור ולברר את התשלום ל-S99: מי אישר? יש חשבונית? זו נורת אזהרה למעילה. להעביר לסמנכ\"ל הכספים מיד.",
                "4) להעביר את שכירות אוגוסט לתקופה הנכונה. 5) לבדוק מול מכירות את החשבונית של 450 אלף.",
                "",
                "בקרות מוצעות: הפרדת תפקידים (מי מקים ספק לא מאשר תשלום), אישור כפול לתשלום מעל 20 אלף,",
                "דוח חריגים שבועי אוטומטי (כפילויות, סופ\"ש, ספק חדש, סכום עגול).",
            ]
            for i, t in enumerate(txt, start=3):
                ws4.cell(row=i, column=1, value=t).fill = SOL_FILL if t else PatternFill()
        else:
            for i in range(3, 20):
                ws4.cell(row=i, column=1).fill = INPUT_FILL
        wb.save((SOL if solution else ROOT) / ("L11_capstone_monthly_report" + ("_solution" if solution else "") + ".xlsx"))


if __name__ == "__main__":
    build_01()
    build_02()
    build_03()
    build_04()
    build_05()
    build_06()
    build_07()
    build_08()
    build_09()
    for p in sorted(ROOT.rglob("*.xlsx")):
        print(p.relative_to(ROOT.parent), p.stat().st_size)
